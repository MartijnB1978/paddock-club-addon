"""Import a season from the original Excel workbook.

Imported weekends keep exactly the points the Excel gave (they are frozen),
predictions and results are imported for viewing.
"""
import difflib
import json
import re
import unicodedata

import openpyxl

from . import db

EMPTY = {"", "nog niet ingevuld", "none"}
ALIASES = {
    "bahrein": "bahrain", "saudi": "saudi arabian", "australie": "australian", "australia": "australian",
    "japan": "japanese", "china": "chinese", "imola": "emilia romagna", "spanje": "spanish",
    "spain": "spanish", "oostenrijk": "austrian", "austria": "austrian", "great brittain": "british",
    "britain": "british", "brittain": "british", "hungary": "hungarian", "belgium": "belgian",
    "netherlands": "dutch", "nederland": "dutch", "zandvoort": "dutch", "monza": "italian",
    "italy": "italian", "canada": "canadian", "mexico": "mexico city", "brazil": "sao paulo",
    "brazilie": "sao paulo", "usa 1": "miami", "usa 2": "united states", "usa 3": "las vegas",
    "esp 1": "barcelona", "barcelona catalunya": "barcelona", "uae": "abu dhabi", "mco": "monaco",
    "malaysian": "bahrain", "malaysia": "bahrain", "great britain": "british", "emilia romagna": "emilia romagna",
    "canadian": "canadian", "dutch": "dutch", "belgian": "belgian", "austrian": "austrian", "verenigde staten": "united states", "usa": "united states",
    "azerbijan": "azerbaijan", "abu dhabi": "abu dhabi", "las vegas": "las vegas",
}


def _plain(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()


def _val(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if (s.lower() in EMPTY or s.lower().startswith("nog niet")) else s


def _label_key(label):
    p = _plain(label).replace("gp", "").replace("grand prix", "").replace("resultaten", "") \
        .replace("resulaten", "").replace("sprint", "").replace("race", "").strip()
    p = re.sub(r"\s+\d+$", "", p).strip()
    for k, v in sorted(ALIASES.items(), key=lambda kv: -len(kv[0])):
        if k in p:
            return v
    return p


def _similar(label, ev):
    key = _label_key(label)
    best = 0.0
    for cand in (ev["name"], ev["country"], ev["location"]):
        c = _plain(cand).replace("grand prix", "").strip()
        if key and (key in c or c in key):
            return 1.0
        best = max(best, difflib.SequenceMatcher(None, key, c).ratio())
    return best


def _map_labels(labels, events):
    """Map Excel race labels onto calendar events: by name first, by calendar order for the rest."""
    events = sorted(events, key=lambda e: e["round"])
    out = [None] * len(labels)
    used = set()
    keys = [_label_key(lab) for lab in labels]
    for i, lab in enumerate(labels):
        if keys.count(keys[i]) > 1:
            continue        # e.g. "Italy GP" and "Italy GP 2": decide by order
        best, bs = None, 0.0
        for e in events:
            if e["id"] in used:
                continue
            sc = _similar(lab, e)
            if sc > bs:
                best, bs = e, sc
        if best and bs >= 0.9:
            out[i] = best
            used.add(best["id"])
    for i in range(len(labels)):
        if out[i]:
            continue
        lo = max([out[j]["round"] for j in range(i) if out[j]] or [0])
        hi = min([out[j]["round"] for j in range(i + 1, len(labels)) if out[j]] or [10_000])
        cands = [e for e in events if e["id"] not in used and lo < e["round"] < hi]
        if cands:
            out[i] = cands[0]
            used.add(cands[0]["id"])
    return out


def _reset_import(c, year):
    """Undo an earlier import of this season so a re-import starts clean."""
    for (eid,) in c.execute("SELECT id FROM events WHERE year=? AND imported=1", (year,)).fetchall():
        c.execute("DELETE FROM scores WHERE event_id=?", (eid,))
        c.execute("DELETE FROM results WHERE event_id=? AND source='import'", (eid,))
        c.execute("DELETE FROM predictions WHERE event_id=?", (eid,))
        c.execute("DELETE FROM entrants WHERE event_id=? AND colour=''", (eid,))
        c.execute("UPDATE sessions SET status='scheduled' WHERE event_id=? AND status='final'", (eid,))
        c.execute("UPDATE events SET imported=0, frozen=0 WHERE id=?", (eid,))
    c.execute("DELETE FROM season_scores WHERE year=?", (year,))


class Drivers:
    def __init__(self, pairs, entrants):
        self.by_name = {}
        for name, tla in pairs:
            self.by_name[_plain(name)] = tla
        for e in entrants:
            self.by_name.setdefault(_plain(e["name"]), e["tla"])
        self.tlas = set(self.by_name.values())

    def tla(self, name):
        name = _val(name)
        if not name:
            return None
        p = _plain(name)
        if p in self.by_name:
            return self.by_name[p]
        if name.upper() in self.tlas:
            return name.upper()
        m = difflib.get_close_matches(p, self.by_name.keys(), n=1, cutoff=0.8)
        if m:
            return self.by_name[m[0]]
        last = p.split(" ")[-1]
        hits = [t for n, t in self.by_name.items() if n.split(" ")[-1] == last]
        return hits[0] if len(set(hits)) == 1 else name


def _team(name, teams):
    name = _val(name)
    if not name or not teams:
        return name
    p = _plain(name)
    plain = {_plain(t): t for t in teams}
    if p in plain:
        return plain[p]
    for k, t in plain.items():
        if p and (p in k or k in p):
            return t
    m = difflib.get_close_matches(p, plain.keys(), n=1, cutoff=0.6)
    return plain[m[0]] if m else name


def _find_header(ws, text, row=1, start=1):
    for c in range(start, ws.max_column + 1):
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and text in v.lower():
            return c
    return None


def parse_result_sheet(ws, drv, teams):
    res = {}
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if isinstance(v, str):
            hdr.setdefault(v.strip().lower(), []).append(c)
    cour = hdr.get("coureurs", [])
    if cour:
        res["race_pos"] = {i - 2: drv.tla(ws.cell(row=i, column=cour[0]).value) for i in range(3, 23)}
    sprint_col = next((c for c in cour if c > (hdr.get("sprint uitslag", [99])[0])), None)
    if sprint_col:
        res["sprint_pos"] = {i - 2: drv.tla(ws.cell(row=i, column=sprint_col).value) for i in range(3, 11)}
    single = {"kwalificatie pole position": "pole", "driver of the day": "dotd", "fastest lap": "fl",
              "kwalificatie pole sprint": "sprint_pole", "fastest pitstop": "fp", "fp": "fp"}
    for h, f in single.items():
        if h in hdr:
            v = ws.cell(row=3, column=hdr[h][0]).value
            res[f] = {1: _team(v, teams) if f == "fp" else drv.tla(v)}
    return {f: {k: v for k, v in slots.items() if v} for f, slots in res.items() if any(slots.values())}


def parse_player(ws, drv, teams):
    """-> season {slot: tla}, races [ {field: {slot: value}} ] in sheet order"""
    season, races = {}, []
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if not isinstance(h, str) or "voorspelling" not in h.lower():
            continue
        if "wk" in h.lower():
            season = {i - 2: drv.tla(ws.cell(row=i, column=c).value) for i in range(3, 13)}
            continue
        p = {"race_pos": {i - 2: drv.tla(ws.cell(row=i, column=c).value) for i in range(3, 13)}}
        for r in range(14, 19):
            lab = str(ws.cell(row=r, column=c - 1).value or "").strip().lower()
            f = {"fl": "fl", "dotd": "dotd", "pole": "pole", "fp": "fp"}.get(lab)
            if f:
                v = ws.cell(row=r, column=c).value
                p[f] = {1: _team(v, teams) if f == "fp" else drv.tla(v)}
        for r in range(18, 24):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "sprint voorspelling" in v.lower():
                sp, spole = {}, None
                for rr in range(r + 1, r + 8):
                    lab = ws.cell(row=rr, column=c - 1).value
                    val = ws.cell(row=rr, column=c).value
                    if isinstance(lab, (int, float)) and 1 <= int(lab) <= 3:
                        sp[int(lab)] = drv.tla(val)
                    elif isinstance(lab, str) and lab.strip().lower() == "pole":
                        spole = drv.tla(val)
                        break
                p["sprint_pos"] = sp
                if spole:
                    p["sprint_pole"] = {1: spole}
                break
        races.append({f: {k: v for k, v in s.items() if v} for f, s in p.items()})
    return {k: v for k, v in season.items() if v}, races


def plan(path, year):
    wb = openpyxl.load_workbook(path, data_only=True)
    if any(n.lower().startswith("players") for n in wb.sheetnames):
        return plan_v2(wb, year)
    events = [dict(r) for r in db.q("SELECT * FROM events WHERE year=? AND cancelled=0 ORDER BY round", year)]
    if not events:
        raise ValueError(f"No {year} calendar available to match the races to")
    dn = wb["Deelnemers"]
    labels, cols = [], []
    for c in range(3, dn.max_column + 1):
        h = dn.cell(row=1, column=c).value
        if not h or "totaal" in str(h).lower():
            break
        labels.append(str(h))
        cols.append(c)
    players = []
    for r in range(3, 40):
        nm = dn.cell(row=r, column=1).value
        if not nm:
            break
        pts = [dn.cell(row=r, column=c).value for c in cols]
        players.append({"name": str(nm).strip(), "season_pts": dn.cell(row=r, column=2).value or 0,
                        "race_pts": pts})
    played = [i for i in range(len(labels)) if any(isinstance(p["race_pts"][i], (int, float)) for p in players)]
    labels = [labels[i] for i in played]
    for p in players:
        p["race_pts"] = [p["race_pts"][i] for i in played]
    mapped = _map_labels(labels, events)

    pairs = []
    if "Teams" in wb.sheetnames:
        t = wb["Teams"]
        for r in range(1, t.max_row + 1):
            if t.cell(row=r, column=4).value and t.cell(row=r, column=5).value:
                pairs.append((t.cell(row=r, column=4).value, t.cell(row=r, column=5).value))
    team_rows = []
    if "Teams" in wb.sheetnames:
        t, cur = wb["Teams"], None
        for r in range(1, t.max_row + 1):
            cur = t.cell(row=r, column=2).value or cur
            if t.cell(row=r, column=5).value:
                team_rows.append({"name": t.cell(row=r, column=4).value, "tla": t.cell(row=r, column=5).value,
                                  "number": str(t.cell(row=r, column=6).value or ""), "team": cur})
    ents = db.q("SELECT x.* FROM entrants x JOIN events e ON e.id=x.event_id WHERE e.year=?", year)
    drv = Drivers(pairs, ents)
    teams = sorted({e["team"] for e in ents if e["team"]} | {r["team"] for r in team_rows if r["team"]})

    res_sheets = [ws for ws in wb.worksheets if "resul" in ws.title.lower()]
    results = [parse_result_sheet(ws, drv, teams) for ws in res_sheets[:len(labels)]]
    preds = {}
    season_preds = {}
    for p in players:
        if p["name"] in wb.sheetnames:
            s, races = parse_player(wb[p["name"]], drv, teams)
            season_preds[p["name"]] = s
            preds[p["name"]] = races[:len(labels)]
    all_events_done = len(labels) >= len(events)
    return {
        "year": year, "labels": labels,
        "events": [{"label": lab, "event_id": ev["id"] if ev else None, "event_name": ev["name"] if ev else "?",
                    "sheet": res_sheets[i].title if i < len(res_sheets) else None,
                    "results": results[i] if i < len(results) else {}}
                   for i, (lab, ev) in enumerate(zip(labels, mapped))],
        "players": players, "preds": preds, "season_preds": season_preds,
        "season_final": all_events_done or year < db.now_utc().year,
        "team_rows": team_rows,
    }


def commit(pl):
    if pl.get("format") == "v2":
        return commit_v2(pl)
    year = pl["year"]
    ts = db.iso(db.now_utc())
    with db.tx() as c:
        _reset_import(c, year)
        final = pl["season_final"]
        c.execute("INSERT OR IGNORE INTO seasons(year) VALUES(?)", (year,))
        c.execute("UPDATE seasons SET rules=?, imported=?, final=? WHERE year=?",
                  ("v1" if final else "v2", 1 if final else 0, 1 if final else 0, year))
        uid = {}
        for p in pl["players"]:
            row = c.execute("SELECT id FROM users WHERE lower(name)=lower(?)", (p["name"],)).fetchone()
            if not row:
                c.execute("INSERT INTO users(name, created_at) VALUES(?,?)", (p["name"], ts))
                row = c.execute("SELECT id FROM users WHERE name=?", (p["name"],)).fetchone()
            uid[p["name"]] = row[0]
        for i, e in enumerate(pl["events"]):
            eid = e["event_id"]
            if not eid:
                continue
            c.execute("UPDATE events SET imported=1, frozen=1 WHERE id=?", (eid,))
            c.execute("DELETE FROM results WHERE event_id=?", (eid,))
            for f, slots in e["results"].items():
                for slot, v in slots.items():
                    c.execute("INSERT INTO results VALUES(?,?,?,?,?,?)", (eid, f, int(slot), v, "import", ts))
            c.execute("UPDATE sessions SET status='final' WHERE event_id=?", (eid,))
            if not c.execute("SELECT 1 FROM entrants WHERE event_id=?", (eid,)).fetchone():
                for t in pl["team_rows"]:
                    c.execute("INSERT OR IGNORE INTO entrants VALUES(?,?,?,?,?,?,?)",
                              (eid, t["number"], t["tla"], t["name"], t["team"], "", ""))
            for p in pl["players"]:
                pts = p["race_pts"][i]
                c.execute("INSERT OR REPLACE INTO scores(user_id,event_id,year,points,detail) VALUES(?,?,?,?,?)",
                          (uid[p["name"]], eid, year, int(pts or 0), json.dumps({"imported": True})))
                c.execute("DELETE FROM predictions WHERE user_id=? AND event_id=?", (uid[p["name"]], eid))
                races = pl["preds"].get(p["name"], [])
                if i < len(races):
                    for f, slots in races[i].items():
                        for slot, v in slots.items():
                            c.execute("INSERT OR REPLACE INTO predictions VALUES(?,?,?,?,?,?,?)",
                                      (uid[p["name"]], year, eid, f, int(slot), v, ts))
        for p in pl["players"]:
            c.execute("DELETE FROM predictions WHERE user_id=? AND year=? AND event_id=0", (uid[p["name"]], year))
            for slot, v in pl["season_preds"].get(p["name"], {}).items():
                c.execute("INSERT INTO predictions VALUES(?,?,?,?,?,?,?)",
                          (uid[p["name"]], year, 0, "season_pos", int(slot), v, ts))
            if final:
                c.execute("INSERT OR REPLACE INTO season_scores VALUES(?,?,?)",
                          (uid[p["name"]], year, int(p["season_pts"] or 0)))
    db.logmsg("info", f"Imported Excel season {year}")


# ---------------------------------------------------------------------------
# 2025+ workbook layout ("Players YYYY" sheet, one block per race per player)
# ---------------------------------------------------------------------------
KNOWN_DRIVERS = {
    "max verstappen": "VER", "sergio perez": "PER", "lewis hamilton": "HAM", "george russell": "RUS",
    "charles leclerc": "LEC", "carlos sainz": "SAI", "lando norris": "NOR", "oscar piastri": "PIA",
    "fernando alonso": "ALO", "lance stroll": "STR", "esteban ocon": "OCO", "pierre gasly": "GAS",
    "alexander albon": "ALB", "alex albon": "ALB", "logan sargeant": "SAR", "valtteri bottas": "BOT",
    "valteri bottas": "BOT", "zhou guanyu": "ZHO", "guanyu zhou": "ZHO", "kevin magnussen": "MAG",
    "nico hulkenberg": "HUL", "yuki tsunoda": "TSU", "daniel ricciardo": "RIC", "daniel riccardo": "RIC",
    "oliver bearman": "BEA", "franco colapinto": "COL", "liam lawson": "LAW", "jack doohan": "DOO",
    "kimi antonelli": "ANT", "andrea kimi antonelli": "ANT", "gabriel bortoleto": "BOR",
    "isack hadjar": "HAD", "arvid lindblad": "LIN",
}


def _base_key(text):
    p = _plain(text)
    for w in ("prediction", "sprint", "gp", "grand prix", " s "):
        p = p.replace(w, " ")
    p = re.sub(r"\s+", " ", p).strip()
    return re.sub(r" s$", "", p).strip()


def _team_pairs(wb):
    pairs = list(KNOWN_DRIVERS.items())
    if "Teams" in wb.sheetnames:
        for row in wb["Teams"].iter_rows(values_only=True):
            cells = [c for c in row if c is not None]
            for i, c in enumerate(cells):
                if isinstance(c, str) and re.fullmatch(r"[A-Za-z]{3}", c.strip()) and i > 0:
                    name = next((x for x in reversed(cells[:i]) if isinstance(x, str) and " " in x.strip()), None)
                    if name and not name.lower().startswith("reserve"):
                        pairs.append((name, c.strip().upper()))
    return pairs


def _driver_teams(wb, drv):
    """{tla: team} from a Teams sheet laid out as  Team | short | Driver | TLA  (2026 workbook)."""
    out = {}
    if "Teams" not in wb.sheetnames:
        return out
    ws = wb["Teams"]
    hdr = None
    for r in range(1, 6):
        vals = [str(ws.cell(row=r, column=c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
        if "teams" in vals and "drivers" in vals:
            hdr = (r, vals.index("teams") + 1, vals.index("drivers") + 1)
            break
    if not hdr:
        return out
    r0, tc, dc = hdr
    team = None
    for r in range(r0 + 1, ws.max_row + 1):
        t = ws.cell(row=r, column=tc).value
        d = ws.cell(row=r, column=dc).value
        if isinstance(t, str) and t.strip():
            team = t.strip()
        if isinstance(d, str) and d.strip() and team:
            out[drv.tla(d)] = team
    return out


def _team_names(wb):
    names = set()
    if "Teams" in wb.sheetnames:
        for row in wb["Teams"].iter_rows(values_only=True):
            for c in row:
                if isinstance(c, str) and c.strip() in {
                        "Alpine", "Aston Martin", "Williams", "Audi", "Cadillac", "Ferrari", "Haas",
                        "McLaren", "Mercedes", "Racing Bulls", "Red Bull Racing", "Kick Sauber", "Sauber"}:
                    names.add(c.strip())
    return names


def _parse_blocks(ws, drv, teams):
    """All prediction blocks of one player sheet, in sheet order."""
    heads = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "prediction" in c.value.lower():
                heads.append((c.column, c.row, c.value))
    heads.sort()
    out = []
    for col, row, text in heads:
        low = text.lower()
        season = "world championship" in low
        mode = "sprint" if "sprint" in low else "race"
        blk = {"name": text, "key": _base_key(text), "season": season, "sprint_header": mode == "sprint",
               "pred": {}, "total": None}
        r = row + 1
        while r <= ws.max_row and r <= row + 30:
            lab = ws.cell(row=r, column=col - 1).value
            val = ws.cell(row=r, column=col).value
            score = ws.cell(row=r, column=col + 1).value
            if isinstance(val, str) and "prediction" in val.lower():
                break
            labs = str(lab).strip().lower() if lab is not None else ""
            vals = str(val).strip().lower() if isinstance(val, str) else ""
            if vals.startswith("total"):
                blk["total"] = score if isinstance(score, (int, float)) else 0
                break
            if season:
                if isinstance(lab, (int, float)) and 1 <= int(lab) <= 10 and r - row <= 12:
                    pass
                if isinstance(val, str) and isinstance(lab, (int, float)) and 1 <= int(lab) <= 10:
                    blk["pred"].setdefault("season_pos", {})[int(lab)] = drv.tla(val)
                elif isinstance(val, str) and lab is None and r - row <= 10 and not vals.startswith("total"):
                    blk["pred"].setdefault("season_pos", {})[r - row] = drv.tla(val)
                if labs.startswith("total") or vals == "total:":
                    blk["total"] = score
                r += 1
                continue
            if labs.startswith("sprint pole"):
                blk["pred"].setdefault("sprint_pole", {})[1] = drv.tla(val)
            elif labs.startswith("sprint"):
                mode = "sprint"
            elif labs.startswith("race"):
                mode = "race"
            elif isinstance(lab, (int, float)) and 1 <= int(lab) <= 10:
                f = "sprint_pos" if mode == "sprint" else "race_pos"
                blk["pred"].setdefault(f, {})[int(lab)] = drv.tla(val)
            elif labs.startswith("pole"):
                f = "sprint_pole" if mode == "sprint" else "pole"
                blk["pred"].setdefault(f, {})[1] = drv.tla(val)
            elif labs.startswith("fastest lap"):
                blk["pred"].setdefault("fl", {})[1] = drv.tla(val)
            elif labs.startswith("driver of the day"):
                blk["pred"].setdefault("dotd", {})[1] = drv.tla(val)
            elif labs.startswith("fastest pit"):
                blk["pred"].setdefault("fp", {})[1] = _team(val, teams)
            r += 1
        if season and blk["total"] is None:
            # WC block: total row directly below the 10 picks
            for rr in range(row + 1, row + 15):
                if str(ws.cell(row=rr, column=col).value or "").strip().lower().startswith("total"):
                    blk["total"] = ws.cell(row=rr, column=col + 1).value
                    break
        blk["pred"] = {f: {k: v for k, v in s.items() if v} for f, s in blk["pred"].items()}
        blk["pred"] = {f: s for f, s in blk["pred"].items() if s}
        out.append(blk)
    return out


def _parse_result_v2(ws, drv, teams):
    hdr = {}
    for r in (1, 6):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                hdr[v.strip().lower()] = (r, c)
    res = {}
    if "driver" in hdr:
        r, c = hdr["driver"]
        res["race_pos"] = {i - r: drv.tla(ws.cell(row=i, column=c).value) for i in range(r + 1, r + 23)}
    single = {"qualification pole": "pole", "driver of the day": "dotd", "fastest lap": "fl",
              "fastest pitstop": "fp", "sprint qualification pole": "sprint_pole", "q. sprint pole": "sprint_pole"}
    for h, f in single.items():
        if h in hdr:
            r, c = hdr[h]
            v = ws.cell(row=r + 1, column=c).value
            res[f] = {1: _team(v, teams) if f == "fp" else drv.tla(v)}
    for h in ("sprint race", "top 8 sprint"):
        if h in hdr:
            r, c = hdr[h]
            res["sprint_pos"] = {i - r: drv.tla(ws.cell(row=i, column=c).value) for i in range(r + 1, r + 9)}
    return {f: {k: v for k, v in s.items() if v} for f, s in res.items() if any(s.values())}


def plan_v2(wb, year):
    events = [dict(r) for r in db.q("SELECT * FROM events WHERE year=? AND cancelled=0 ORDER BY round", year)]
    if not events:
        raise ValueError(f"No {year} calendar available to match the races to")
    psheet = next(n for n in wb.sheetnames if n.lower().startswith("players"))
    ps = wb[psheet]
    hdr_row = next(r for r in range(1, 40) if str(ps.cell(row=r, column=2).value or "").strip().lower() == "player")
    players = []
    excel_total = {}
    for r in range(hdr_row + 1, hdr_row + 30):
        nm = ps.cell(row=r, column=2).value
        if not nm:
            if players:
                break
            continue
        nm = str(nm).strip()
        if nm in wb.sheetnames:
            players.append(nm)
            excel_total[nm] = ps.cell(row=r, column=3).value
    pairs = _team_pairs(wb)
    ents = db.q("SELECT x.* FROM entrants x JOIN events e ON e.id=x.event_id WHERE e.year=?", year)
    drv = Drivers(pairs, ents)
    teams = sorted(_team_names(wb) | {e["team"] for e in ents if e["team"]})

    # result sheets by base name
    results = {}
    for ws in wb.worksheets:
        if ws.title in players or ws.title == psheet:
            continue
        if ws.cell(row=1, column=2).value and str(ws.cell(row=1, column=2).value).strip().lower() == "driver":
            results[_base_key(ws.title)] = (ws.title, _parse_result_v2(ws, drv, teams))

    # blocks per player; race order from the first player's sheet
    blocks = {p: _parse_blocks(wb[p], drv, teams) for p in players}
    order = []
    for p in players:
        for b in blocks[p]:
            if not b["season"] and b["key"] not in order:
                order.append(b["key"])
        if len(order) > 5:
            break
    # weekends that took place (a race result exists) or are still to come (in calendar)
    raced = [k for k in order if results.get(k, ("", {}))[1].get("race_pos")]
    keys = [k for k in order if k in raced or not k.startswith(("bahrein", "bahrain", "saudi"))]
    mapped = dict(zip(keys, _map_labels(keys, events)))

    plan_events = []
    for k in order:
        ev = mapped.get(k)
        if not ev:
            continue
        sheet, res = results.get(k, (None, {}))
        played = bool(res.get("race_pos"))
        pts, preds = {}, {}
        for p in players:
            tot, merged = None, {}
            for b in blocks[p]:
                if b["key"] != k or b["season"]:
                    continue
                for f, s in b["pred"].items():
                    merged.setdefault(f, {}).update(s)
                if b["total"] is not None and not b["sprint_header"]:
                    tot = b["total"]
            pts[p] = int(tot or 0)
            preds[p] = merged
        plan_events.append({"label": k.title(), "event_id": ev["id"], "event_name": ev["name"], "sheet": sheet,
                            "played": played, "results": res if played else {}, "points": pts, "preds": preds})
    season_preds, season_pts = {}, {}
    for p in players:
        sb = next((b for b in blocks[p] if b["season"]), None)
        season_preds[p] = sb["pred"].get("season_pos", {}) if sb else {}
        season_pts[p] = int((sb or {}).get("total") or 0)
    standings = {}
    if "Drivers Championship" in wb.sheetnames:
        dc = wb["Drivers Championship"]
        for r in range(1, dc.max_row + 1):
            pos, nm = dc.cell(row=r, column=1).value, dc.cell(row=r, column=2).value
            if isinstance(pos, (int, float)) and isinstance(nm, str) and r > 15:
                if int(pos) in standings:
                    break           # end of the drivers table
                standings[int(pos)] = (drv.tla(nm), dc.cell(row=r, column=4).value)
    dteams = _driver_teams(wb, drv)
    future_left = any(not e["played"] for e in plan_events) or len([e for e in plan_events if e["played"]]) < len(events)
    final = year < db.now_utc().year or not future_left
    return {"format": "v2", "year": year, "events": plan_events, "players": [
        {"name": p, "season_pts": season_pts[p], "excel_total": excel_total.get(p),
         "race_pts": [e["points"][p] for e in plan_events if e["played"]]} for p in players],
        "season_preds": season_preds, "season_final": final, "standings": standings,
        "team_rows": [{"name": n, "tla": t, "number": "", "team": dteams.get(t, "")} for n, t in pairs if " " in n]}


def commit_v2(pl):
    year = pl["year"]
    ts = db.iso(db.now_utc())
    final = pl["season_final"]
    with db.tx() as c:
        _reset_import(c, year)
        c.execute("INSERT OR IGNORE INTO seasons(year) VALUES(?)", (year,))
        c.execute("UPDATE seasons SET rules='v2', imported=?, final=? WHERE year=?",
                  (1 if final else 0, 1 if final else 0, year))
        uid = {}
        for p in pl["players"]:
            row = c.execute("SELECT id FROM users WHERE lower(name)=lower(?)", (p["name"],)).fetchone()
            if not row:
                c.execute("INSERT INTO users(name, created_at) VALUES(?,?)", (p["name"], ts))
                row = c.execute("SELECT id FROM users WHERE name=?", (p["name"],)).fetchone()
            uid[p["name"]] = row[0]
        for e in pl["events"]:
            eid = e["event_id"]
            if e["played"]:
                c.execute("UPDATE events SET imported=1, frozen=1 WHERE id=?", (eid,))
                c.execute("DELETE FROM results WHERE event_id=?", (eid,))
                for f, slots in e["results"].items():
                    for slot, v in slots.items():
                        c.execute("INSERT INTO results VALUES(?,?,?,?,?,?)", (eid, f, int(slot), v, "import", ts))
                c.execute("UPDATE sessions SET status='final' WHERE event_id=?", (eid,))
                if not c.execute("SELECT 1 FROM entrants WHERE event_id=?", (eid,)).fetchone():
                    for t in pl["team_rows"]:
                        c.execute("INSERT OR IGNORE INTO entrants VALUES(?,?,?,?,?,?,?)",
                                  (eid, t["number"], t["tla"], t["name"].title(), t["team"], "", ""))
            for p in pl["players"]:
                u = uid[p["name"]]
                if e["played"]:
                    c.execute("INSERT OR REPLACE INTO scores(user_id,event_id,year,points,detail) VALUES(?,?,?,?,?)",
                              (u, eid, year, e["points"][p["name"]], json.dumps({"imported": True})))
                pred = e["preds"].get(p["name"], {})
                if not e["played"] and c.execute("SELECT 1 FROM predictions WHERE user_id=? AND event_id=? "
                                                 "LIMIT 1", (u, eid)).fetchone():
                    continue    # never overwrite predictions made in the app for an open weekend
                if e["played"] or pred:
                    c.execute("DELETE FROM predictions WHERE user_id=? AND event_id=?", (u, eid))
                for f, slots in pred.items():
                    for slot, v in slots.items():
                        c.execute("INSERT OR REPLACE INTO predictions VALUES(?,?,?,?,?,?,?)",
                                  (u, year, eid, f, int(slot), v, ts))
        for p in pl["players"]:
            u = uid[p["name"]]
            c.execute("DELETE FROM predictions WHERE user_id=? AND year=? AND event_id=0", (u, year))
            for slot, v in pl["season_preds"].get(p["name"], {}).items():
                c.execute("INSERT INTO predictions VALUES(?,?,?,?,?,?,?)", (u, year, 0, "season_pos", int(slot), v, ts))
            if final:
                c.execute("INSERT OR REPLACE INTO season_scores VALUES(?,?,?)", (u, year, p["season_pts"]))
        if pl["standings"] and not final:
            c.execute("DELETE FROM standings WHERE year=?", (year,))
            for pos, (tla, pts) in pl["standings"].items():
                c.execute("INSERT OR REPLACE INTO standings VALUES(?,?,?,?,?)", (year, pos, tla, pts, ts))
    db.logmsg("info", f"Imported Excel season {year}")
